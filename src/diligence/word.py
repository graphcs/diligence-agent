"""Render the memo as a Word document with its charts embedded in their sections.

This is the document posted to Slack: Markdown can't carry images inside the file, a .docx
can, and the deal team can edit and comment on it. Same section order as render.to_markdown."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import charts, render
from .ledger import Ledger
from .models import MemoState

FONT = "Arial"  # installed everywhere Word, Google Docs and Slack previews render
INK, INK_2, MUTED, ACCENT, WARN = (RGBColor.from_string(c.lstrip("#").upper()) for c in
                                   (charts.INK, charts.INK_2, charts.MUTED, charts.COMPANY, "#C2491D"))
BANNER_BG = {"draft": "FDF1E6", "in_review": "EEF4FC", "signed_off": "E9F6EF"}
CHART_WIDTH = {"snapshot": 4.6, "product": 6.3, "competitors": 4.0, "financials": 4.4}  # inches
_XML_UNSAFE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _t(text: str) -> str:
    """Model-written or source text minus control characters, which aren't valid in XML."""
    return _XML_UNSAFE.sub("", text)


def _shade(paragraph, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def _field(paragraph, instr: str) -> None:
    """A Word field such as PAGE, updated by Word when the document is opened or printed."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def _styles(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = FONT, Pt(10.5), INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(4)
    for name, size, before in (("Title", 22, 0), ("Heading 1", 14, 14)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = FONT, Pt(size), True, INK
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(4)
        st.paragraph_format.keep_with_next = True
    # The default Title style has a bottom border; drop it for a cleaner header.
    pPr = doc.styles["Title"].element.pPr
    if pPr is not None:
        for b in pPr.findall(qn("w:pBdr")):
            pPr.remove(b)
    lb = doc.styles["List Bullet"]
    lb.font.name, lb.font.size = FONT, Pt(10.5)
    lb.paragraph_format.space_after = Pt(2)


def _small(doc, text: str, italic: bool = False, color=MUTED, size: float = 8.5, after: float = 4):
    p = doc.add_paragraph()
    r = p.add_run(_t(text))
    r.font.size, r.font.italic, r.font.color.rgb = Pt(size), italic, color
    p.paragraph_format.space_after = Pt(after)
    return p


def _bullet(doc, text: str, refs: str = "", tag: str = "") -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(_t(text))
    if refs:
        r = p.add_run(f" {refs}")
        r.font.color.rgb = INK_2
    if tag:
        r = p.add_run(f" ({_t(tag)})")
        r.font.italic, r.font.color.rgb = True, WARN


def to_docx(ledger: Ledger, memo: MemoState, images: dict[str, bytes] | None = None) -> bytes:
    """images: chart filename -> PNG bytes (as written by charts.render_all); drawn when missing."""
    images = charts.render_all(memo) if images is None else images
    cites = render.Citations(ledger, memo)
    doc = Document()
    _styles(doc)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    doc.core_properties.title = f"{render.stage_title(memo)}: {memo.company}"
    doc.core_properties.author = "Diligence agent"

    # Header block: title, meta line, review banner.
    doc.add_paragraph(_t(f"{render.stage_title(memo)}: {memo.company}"), style="Title")
    _small(doc, f"Requested by {memo.requested_by} · version {memo.version} · CRM {memo.deal_id}",
           color=INK_2, size=9, after=8)
    banner = doc.add_paragraph()
    r = banner.add_run(_t(render.banner(memo)))
    r.font.bold, r.font.size = True, Pt(9.5)
    _shade(banner, BANNER_BG[render.review_state(memo)])
    banner.paragraph_format.space_after = Pt(10)

    if memo.summary:
        doc.add_heading("Summary", level=1)
        for p in memo.summary:
            _bullet(doc, p.text, cites.refs(p.evidence_ids))

    drawn = {c.section: c for c in charts.available(memo) if c.filename in images}
    for key, s in memo.sections.items():
        doc.add_heading(_t(s.title), level=1)
        if not s.claims:
            _small(doc, "No supported findings.", italic=True, color=INK_2, size=10.5)
        for c in s.claims:
            _bullet(doc, c.text, cites.refs(c.evidence_ids), render.claim_tag(c))
        if key in drawn:
            c = drawn[key]
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_before = Pt(8)
            pic.paragraph_format.keep_with_next = True
            pic.add_run().add_picture(io.BytesIO(images[c.filename]), width=Inches(CHART_WIDTH[key]))
            cap = _small(doc, f"{c.caption} · sources {cites.refs(charts.evidence_ids(memo, key))}",
                         italic=True, color=INK_2, size=9)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if memo.figures:
        doc.add_heading("Figures to verify", level=1)
        doc.add_paragraph("Every financial figure must be checked against its source by a person on the "
                          "deal team before this memo can be signed off.")
        mark = {"pending": "☐", "verified": "☑", "disputed": "⚠"}
        for f in memo.figures:
            p = doc.add_paragraph()
            p.add_run(f"{mark[f.status]}  ")
            p.add_run(_t(f.id)).font.bold = True
            p.add_run(f" {_t(f.text)} ")
            p.add_run(cites.refs(f.evidence_ids)).font.color.rgb = INK_2
            if f.by:
                p.add_run(_t(render.figure_who(f))).font.italic = True

    if memo.consistency_issues:
        doc.add_heading("Consistency checks", level=1)
        for p in memo.consistency_issues:
            _bullet(doc, p.text, cites.refs(p.evidence_ids))

    doc.add_heading("Gaps and unverified", level=1)
    for g in render.gaps(memo) or ["None reported."]:
        _bullet(doc, g)

    doc.add_heading("Sources", level=1)
    for n, r in cites.sources():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Inches(0.3), Inches(-0.3)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"{n}.\t")
        p.add_run(_t(r.title)).font.bold = True
        where = f", p. {r.page}" if r.page else ""
        as_of = f" · as of {r.as_of}" if r.as_of else ""
        p.add_run(_t(f"{where} ({render.CONF_LABEL[r.confidence]}{as_of}) - "))
        p.add_run(_t(r.source_uri)).font.color.rgb = INK_2
        q = _small(doc, f"“{r.quote}”", italic=True, color=INK_2, size=9.5, after=6)
        q.paragraph_format.left_indent = Inches(0.3)

    if memo.changelog:
        doc.add_heading("Changelog", level=1)
        for c in memo.changelog:
            _bullet(doc, c)

    footer = sec.footer.paragraphs[0]
    run = footer.add_run(_t(f"{memo.company} · v{memo.version} · {render.status_line(memo)}   ·   Page "))
    run.font.size, run.font.color.rgb = Pt(8), MUTED
    _field(footer, "PAGE")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

